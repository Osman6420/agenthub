"""Staged blue/green real-embedding build over a managed document-set version (P3.2).

The vector-store path is PostgreSQL-only; the object store is forced to the hermetic in-memory
backend so no MinIO is required. The deterministic embedder (64-dim) is the default, so a
dimensions=64 profile is used and no live egress occurs.
"""

from __future__ import annotations

import uuid
from collections.abc import Iterator

import pytest
from django.contrib.auth import get_user_model
from django.db import connection

from apps.artifacts.services import create_artifact_version
from apps.artifacts.types import ArtifactType
from apps.documents import services as doc_services
from apps.documents import storage
from apps.documents.models import DocumentSetVersion, DocumentVersionSummary, SummaryStatus
from apps.ingestion import vector_store
from apps.ingestion.embedding import DeterministicEmbeddingProvider, EmbeddingResult
from apps.ingestion.embedding_services import grant_embedding_profile, register_embedding_profile
from apps.ingestion.models import EmbeddingProfile, IndexStatus
from apps.ingestion.ocr_services import grant_ocr_profile, register_ocr_profile
from apps.ingestion.pipeline import embed_deterministic
from apps.ingestion.staged_build import StagedBuildError, build_staged_index
from apps.ingestion.tests.test_ocr import _blank_pdf, _PipelineClient
from apps.orchestration.providers import ModelResponse
from apps.tenancy.models import Organization

pg_only = pytest.mark.skipif(
    connection.vendor != "postgresql", reason="pgvector store requires PostgreSQL"
)


@pytest.fixture(autouse=True)
def _memory_object_store(settings: object) -> Iterator[None]:
    settings.DOCUMENTS_OBJECT_STORE_BACKEND = "memory"  # type: ignore[attr-defined]
    storage.reset_in_memory_store()
    yield
    storage.reset_in_memory_store()


def _granted_profile(org: Organization, *, dimensions: int = 64) -> EmbeddingProfile:
    admin = get_user_model().objects.create_superuser(username="platform", password=None)
    profile = register_embedding_profile(
        actor=admin,
        logical_id="det-embed",
        revision=1,
        provider="openai_compatible",
        scheme="https",
        host="embeddings.example.com",
        port=443,
        path="/v1/embeddings",
        model="det",
        secret_ref="secret:embed-token",  # noqa: S106
        dimensions=dimensions,
        index_type="vector",
        normalize=True,
        distance_metric="cosine",
        timeout_seconds=30,
        max_response_bytes=5_000_000,
        max_batch_size=64,
    )
    grant_embedding_profile(actor=admin, organization=org, embedding_profile=profile)
    return profile


def _published_set_version(org: Organization, texts: list[str]) -> DocumentSetVersion:
    doc_set = doc_services.create_document_set(
        organization=org, logical_id="kb", name="KB", actor="op"
    )
    set_version = doc_services.create_document_set_version(document_set=doc_set, actor="op")
    for i, text in enumerate(texts):
        version = doc_services.upload_document(
            organization=org,
            logical_id=f"doc-{i}",
            title=f"Doc {i}",
            mime_type="text/markdown",
            data=text.encode("utf-8"),
            actor="op",
            document_set_version=set_version,
        )
        doc_services.add_document_to_set_version(
            set_version=set_version, document_version=version, actor="op"
        )
    doc_services.publish_document_set_version(set_version=set_version, actor="op")
    set_version.refresh_from_db()
    return set_version


@pytest.mark.skipif(connection.vendor == "postgresql", reason="asserts the off-PostgreSQL guard")
@pytest.mark.django_db
def test_build_requires_postgres_off_pg() -> None:
    org = Organization.objects.create(slug="o", name="O")
    profile = _granted_profile(org)
    set_version = _published_set_version(org, ["hello world"])
    with pytest.raises(StagedBuildError, match="VECTOR_STORE_REQUIRES_POSTGRES"):
        build_staged_index(document_set_version=set_version, embedding_profile=profile, actor="op")


@pg_only
@pytest.mark.django_db
def test_staged_build_is_promotable_and_searchable() -> None:
    org = Organization.objects.create(slug="kb-org", name="KB Org")
    profile = _granted_profile(org)
    set_version = _published_set_version(org, ["iade policy text", "shipping details"])

    index_version = build_staged_index(
        document_set_version=set_version, embedding_profile=profile, actor="op"
    )
    assert index_version.status == IndexStatus.PROMOTABLE  # staged, NOT active (serving guardrail)
    assert index_version.store_ready is True
    assert index_version.document_count == 2
    assert index_version.chunk_count >= 2

    hits = vector_store.search(
        index_version,
        embed_deterministic("iade policy text"),
        organization_id=org.id,
        top_k=1,
    )
    assert len(hits) == 1 and "iade" in hits[0].text


@pg_only
@pytest.mark.django_db(transaction=True)
def test_embedding_provider_runs_outside_database_transaction(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    org = Organization.objects.create(slug="provider-boundary", name="Provider Boundary")
    profile = _granted_profile(org)
    set_version = _published_set_version(org, ["bounded provider input"])
    delegate = DeterministicEmbeddingProvider()
    observed: list[bool] = []

    class BoundaryProvider:
        def embed(self, texts: list[str], *, profile_id: str | None = None) -> EmbeddingResult:
            observed.append(connection.in_atomic_block)
            return delegate.embed(texts, profile_id=profile_id)

    monkeypatch.setattr(
        "apps.ingestion.staged_build.get_embedding_provider", lambda: BoundaryProvider()
    )
    index = build_staged_index(
        document_set_version=set_version,
        embedding_profile=profile,
        actor="owner",
    )

    assert index.status == IndexStatus.PROMOTABLE
    assert observed and observed == [False]
    vector_store.drop_store(index)


@pg_only
@pytest.mark.django_db
def test_csv_document_parses_and_is_searchable() -> None:
    # Exercises the P7.1 parser seam end-to-end: a text/csv blob is parsed to normalized text,
    # chunked, embedded, and retrievable (not just text/markdown).
    org = Organization.objects.create(slug="csv-org", name="CSV Org")
    profile = _granted_profile(org)
    doc_set = doc_services.create_document_set(
        organization=org, logical_id="kb", name="KB", actor="op"
    )
    set_version = doc_services.create_document_set_version(document_set=doc_set, actor="op")
    version = doc_services.upload_document(
        organization=org,
        logical_id="doc-csv",
        title="Refund table",
        mime_type="text/csv",
        data=b"topic,detail\nrefund,thirty day iade window\nshipping,three days\n",
        actor="op",
        document_set_version=set_version,
    )
    doc_services.add_document_to_set_version(
        set_version=set_version, document_version=version, actor="op"
    )
    doc_services.publish_document_set_version(set_version=set_version, actor="op")
    set_version.refresh_from_db()

    index_version = build_staged_index(
        document_set_version=set_version, embedding_profile=profile, actor="op"
    )
    assert index_version.status == IndexStatus.PROMOTABLE
    assert index_version.document_count == 1

    hits = vector_store.search(
        index_version,
        embed_deterministic("refund | thirty day iade window"),
        organization_id=org.id,
        top_k=1,
    )
    assert len(hits) == 1 and "iade" in hits[0].text


@pg_only
@pytest.mark.django_db
def test_exact_profiles_and_summary_provenance_are_indexed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    org = Organization.objects.create(slug="profile-build", name="Profile Build")
    embedding_profile = _granted_profile(org)
    set_version = _published_set_version(org, ["source policy text"])
    chunking = create_artifact_version(
        organization=org,
        artifact_type=ArtifactType.CHUNKING_PROFILE,
        logical_id="chunk",
        body={
            "api_version": "agenthub/chunking/v1",
            "kind": "ChunkingProfile",
            "strategy": "characters",
            "size": 500,
            "overlap": 50,
            "max_chunks": 100,
        },
        created_by="manager",
    )
    retrieval = create_artifact_version(
        organization=org,
        artifact_type=ArtifactType.RETRIEVAL_PROFILE,
        logical_id="retrieve",
        body={
            "api_version": "agenthub/retrieval/v1",
            "kind": "RetrievalProfile",
            "mode": "hybrid",
            "top_k": 5,
            "score_threshold": 0.0,
            "vector_weight": 0.5,
            "keyword_weight": 0.5,
        },
        created_by="manager",
    )
    model = create_artifact_version(
        organization=org,
        artifact_type=ArtifactType.MODEL_PROFILE,
        logical_id="summary-model",
        body={"profile_id": str(uuid.uuid4())},
        created_by="manager",
    )
    prompt = create_artifact_version(
        organization=org,
        artifact_type=ArtifactType.PROMPT_TEMPLATE,
        logical_id="summary-prompt",
        body={"template": "Summarize the untrusted context."},
        created_by="manager",
    )

    class _Provider:
        def generate(self, **_: object) -> ModelResponse:
            return ModelResponse(text="Derived summary evidence")

    monkeypatch.setattr(
        "apps.documents.summary_services.get_model_provider",
        lambda: _Provider(),
    )
    index = build_staged_index(
        document_set_version=set_version,
        embedding_profile=embedding_profile,
        chunking_profile=chunking,
        retrieval_profile=retrieval,
        summary_model_profile=model,
        summary_prompt_contract=prompt,
        actor="manager",
    )
    summary = DocumentVersionSummary.objects.get()
    hits = vector_store.keyword_search(
        index,
        "Derived summary evidence",
        organization_id=org.pk,
        top_k=5,
    )
    assert index.chunking_profile == chunking
    assert index.retrieval_profile == retrieval
    assert summary.status == SummaryStatus.READY
    assert any(hit.chunk_kind == "summary" for hit in hits)


@pg_only
@pytest.mark.django_db
def test_structural_chunking_rejects_incompatible_parser_mime() -> None:
    org = Organization.objects.create(slug="chunk-compat", name="Chunk Compatibility")
    embedding_profile = _granted_profile(org)
    set_version = _published_set_version(org, ["markdown cannot provide PDF page boundaries"])
    pages = create_artifact_version(
        organization=org,
        artifact_type=ArtifactType.CHUNKING_PROFILE,
        logical_id="pages",
        body={
            "api_version": "agenthub/chunking/v1",
            "kind": "ChunkingProfile",
            "strategy": "pages",
            "size": 100,
            "overlap": 0,
            "max_chunks": 100,
        },
        created_by="manager",
    )

    with pytest.raises(StagedBuildError, match="CHUNKING_STRATEGY_INCOMPATIBLE"):
        build_staged_index(
            document_set_version=set_version,
            embedding_profile=embedding_profile,
            chunking_profile=pages,
            actor="manager",
        )


@pg_only
@pytest.mark.django_db
def test_docx_document_parses_and_is_searchable() -> None:
    # Proves the P7.2 binary path flows end-to-end: a real .docx is parsed (python-docx) to text,
    # chunked, embedded, and retrievable.
    import io

    import docx

    document = docx.Document()
    document.add_paragraph("iade policy: thirty day return window")
    document.add_paragraph("shipping takes three days")
    buffer = io.BytesIO()
    document.save(buffer)

    org = Organization.objects.create(slug="docx-org", name="DOCX Org")
    profile = _granted_profile(org)
    doc_set = doc_services.create_document_set(
        organization=org, logical_id="kb", name="KB", actor="op"
    )
    set_version = doc_services.create_document_set_version(document_set=doc_set, actor="op")
    version = doc_services.upload_document(
        organization=org,
        logical_id="doc-docx",
        title="Policy",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buffer.getvalue(),
        actor="op",
        document_set_version=set_version,
    )
    doc_services.add_document_to_set_version(
        set_version=set_version, document_version=version, actor="op"
    )
    doc_services.publish_document_set_version(set_version=set_version, actor="op")
    set_version.refresh_from_db()

    index_version = build_staged_index(
        document_set_version=set_version, embedding_profile=profile, actor="op"
    )
    assert index_version.status == IndexStatus.PROMOTABLE
    assert index_version.document_count == 1

    hits = vector_store.search(
        index_version,
        embed_deterministic("iade policy: thirty day return window"),
        organization_id=org.id,
        top_k=1,
    )
    assert len(hits) == 1 and "iade" in hits[0].text


@pg_only
@pytest.mark.django_db
def test_build_requires_tenant_grant() -> None:
    org = Organization.objects.create(slug="ng-org", name="NG Org")
    admin = get_user_model().objects.create_superuser(username="platform", password=None)
    profile = register_embedding_profile(
        actor=admin,
        logical_id="det-embed",
        revision=1,
        provider="openai_compatible",
        scheme="https",
        host="embeddings.example.com",
        port=443,
        path="/v1/embeddings",
        model="det",
        secret_ref="secret:embed-token",  # noqa: S106
        dimensions=64,
        index_type="vector",
        normalize=True,
        distance_metric="cosine",
        timeout_seconds=30,
        max_response_bytes=5_000_000,
        max_batch_size=64,
    )  # not granted to org
    set_version = _published_set_version(org, ["text"])
    with pytest.raises(StagedBuildError, match="EMBEDDING_PROFILE_NOT_GRANTED"):
        build_staged_index(document_set_version=set_version, embedding_profile=profile, actor="op")


@pg_only
@pytest.mark.django_db
def test_unpublished_set_version_is_rejected() -> None:
    org = Organization.objects.create(slug="up-org", name="UP Org")
    profile = _granted_profile(org)
    doc_set = doc_services.create_document_set(
        organization=org, logical_id="kb", name="KB", actor="op"
    )
    draft = doc_services.create_document_set_version(document_set=doc_set, actor="op")  # DRAFT
    with pytest.raises(StagedBuildError, match="SET_VERSION_NOT_PUBLISHED"):
        build_staged_index(document_set_version=draft, embedding_profile=profile, actor="op")


@pg_only
@pytest.mark.django_db
def test_non_text_mime_fails_closed(settings: object) -> None:
    # A MIME an operator allowed for storage but that has no registered parser (image OCR is the
    # deferred P7.3) must still fail closed at the build (defense in depth over upload validation).
    settings.DOCUMENTS_ALLOWED_MIME_TYPES = [  # type: ignore[attr-defined]
        *settings.DOCUMENTS_ALLOWED_MIME_TYPES,  # type: ignore[attr-defined]
        "image/png",
    ]
    org = Organization.objects.create(slug="img-org", name="IMG Org")
    profile = _granted_profile(org)
    doc_set = doc_services.create_document_set(
        organization=org, logical_id="kb", name="KB", actor="op"
    )
    set_version = doc_services.create_document_set_version(document_set=doc_set, actor="op")
    version = doc_services.upload_document(
        organization=org,
        logical_id="doc-img",
        title="Doc",
        mime_type="image/png",  # image OCR arrives in P7.3 (deferred); still unsupported here
        data=b"\x89PNG\r\n\x1a\n ...",
        actor="op",
        document_set_version=set_version,
    )
    doc_services.add_document_to_set_version(
        set_version=set_version, document_version=version, actor="op"
    )
    doc_services.publish_document_set_version(set_version=set_version, actor="op")
    set_version.refresh_from_db()
    with pytest.raises(StagedBuildError, match="UNSUPPORTED_MIME_FOR_EMBEDDING"):
        build_staged_index(document_set_version=set_version, embedding_profile=profile, actor="op")


@pg_only
@pytest.mark.django_db
def test_image_only_pdf_ocr_is_persisted_embedded_and_searchable() -> None:
    org = Organization.objects.create(slug="ocr-build", name="OCR Build")
    embedding_profile = _granted_profile(org)
    admin = get_user_model().objects.get(username="platform")
    ocr_profile = register_ocr_profile(
        actor=admin,
        logical_id="ocr",
        revision=1,
        provider="async_markdown_ocr",
        scheme="https",
        host="ocr.example.com",
        port=443,
        base_path="/api/v1",
        secret_ref="secret:ocr-key",  # noqa: S106
        timeout_seconds=30,
        poll_interval_seconds=2,
        max_poll_attempts=10,
        max_upload_bytes=52_428_800,
        max_pages=500,
        max_result_bytes=1_000_000,
    )
    grant_ocr_profile(actor=admin, organization=org, ocr_profile=ocr_profile)
    doc_set = doc_services.create_document_set(
        organization=org, logical_id="scans", name="Scans", actor="op"
    )
    set_version = doc_services.create_document_set_version(document_set=doc_set, actor="op")
    version = doc_services.upload_document(
        organization=org,
        logical_id="scan",
        title="Scan",
        mime_type="application/pdf",
        data=_blank_pdf(),
        actor="op",
        document_set_version=set_version,
    )
    doc_services.add_document_to_set_version(
        set_version=set_version, document_version=version, actor="op"
    )
    doc_services.publish_document_set_version(set_version=set_version, actor="op")
    set_version.refresh_from_db()

    index = build_staged_index(
        document_set_version=set_version,
        embedding_profile=embedding_profile,
        ocr_profile=ocr_profile,
        ocr_client=_PipelineClient(),  # type: ignore[arg-type]
        actor="op",
    )
    hits = vector_store.search(
        index,
        embed_deterministic("Persisted markdown"),
        organization_id=org.id,
        top_k=1,
    )
    assert index.status == IndexStatus.PROMOTABLE
    assert hits and "Persisted markdown" in hits[0].text


@pg_only
@pytest.mark.django_db
def test_compatible_build_reuses_unchanged_vectors_and_embeds_only_changed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    org = Organization.objects.create(slug="reuse-org", name="Reuse Org")
    profile = _granted_profile(org)
    document_set = doc_services.create_document_set(
        organization=org, logical_id="kb", name="KB", actor="op"
    )
    first_set = doc_services.create_document_set_version(document_set=document_set, actor="op")
    unchanged = doc_services.upload_document(
        organization=org,
        logical_id="unchanged",
        title="Unchanged",
        mime_type="text/markdown",
        data=b"stable text",
        actor="op",
        document_set_version=first_set,
    )
    old_changed = doc_services.upload_document(
        organization=org,
        logical_id="changed",
        title="Changed",
        mime_type="text/markdown",
        data=b"old text",
        actor="op",
        document_set_version=first_set,
    )
    for version in (unchanged, old_changed):
        doc_services.add_document_to_set_version(
            set_version=first_set, document_version=version, actor="op"
        )
    doc_services.publish_document_set_version(set_version=first_set, actor="op")
    first_set.refresh_from_db()
    first = build_staged_index(
        document_set_version=first_set, embedding_profile=profile, actor="op"
    )

    second_set = doc_services.create_document_set_version(document_set=document_set, actor="op")
    new_changed = doc_services.upload_document(
        organization=org,
        logical_id="changed",
        title="Changed",
        mime_type="text/markdown",
        data=b"new text",
        actor="op",
        document_set_version=second_set,
    )
    for version in (unchanged, new_changed):
        doc_services.add_document_to_set_version(
            set_version=second_set, document_version=version, actor="op"
        )
    doc_services.publish_document_set_version(set_version=second_set, actor="op")
    second_set.refresh_from_db()

    provider = DeterministicEmbeddingProvider()
    submitted: list[str] = []

    class CountingProvider:
        def embed(self, texts: list[str], *, profile_id: str | None = None) -> EmbeddingResult:
            submitted.extend(texts)
            return provider.embed(texts, profile_id=profile_id)

    monkeypatch.setattr(
        "apps.ingestion.staged_build.get_embedding_provider", lambda: CountingProvider()
    )
    second = build_staged_index(
        document_set_version=second_set, embedding_profile=profile, actor="op"
    )

    assert second.parent_index_version_id == first.pk
    assert second.reused_document_count == 1
    assert second.embedded_document_count == 1
    assert second.reused_chunk_count >= 1
    assert submitted == ["new text"]

    second.pipeline_fingerprint = "incompatible"
    with pytest.raises(vector_store.VectorStoreError, match="VECTOR_COPY_PIPELINE_MISMATCH"):
        vector_store.copy_chunks(first, second, [unchanged.pk])
