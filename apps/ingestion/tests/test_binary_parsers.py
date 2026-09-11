"""Unit tests for the P7.2 local binary-format parsers (pdf/docx/xlsx).

Fixtures are built in-process (docx/xlsx via the same libraries; pdf via a hand-assembled minimal
document with correct xref offsets), so no binary blobs are committed and no network is used.
"""

from __future__ import annotations

import io

import docx
import openpyxl
import pdfplumber
import pytest

from apps.ingestion.parsers import ParserError, parse_document

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _build_minimal_pdf(text: str) -> bytes:
    """Assemble a valid single-page PDF whose content stream draws ``text`` (ASCII, no parens)."""
    stream = b"BT /F1 24 Tf 72 720 Td (" + text.encode("latin-1") + b") Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += str(i).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref_pos = len(pdf)
    size = len(objects) + 1
    pdf += b"xref\n0 " + str(size).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n".encode()
    pdf += b"trailer\n<< /Size " + str(size).encode() + b" /Root 1 0 R >>\n"
    pdf += b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF"
    return bytes(pdf)


def test_pdf_text_is_extracted() -> None:
    out = parse_document("application/pdf", _build_minimal_pdf("iade policy text"))
    assert "iade policy text" in out.text
    assert out.parser == "pdf"
    assert out.page_count == 1


def test_pdf_malformed_fails_closed() -> None:
    with pytest.raises(ParserError) as exc:
        parse_document("application/pdf", b"%PDF-1.4 not really a pdf")
    assert exc.value.code in {"PDF_PARSE_FAILED", "EMPTY_DOCUMENT"}
    # No document bytes leak into the error message.
    assert str(exc.value) == exc.value.code


def test_mixed_pdf_requires_ocr_instead_of_silently_dropping_page(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class _Pdf:
        pages = [
            type("Page", (), {"extract_text": lambda self: "visible text"})(),
            type("Page", (), {"extract_text": lambda self: ""})(),
        ]

        def __enter__(self) -> _Pdf:
            return self

        def __exit__(self, *args: object) -> None:
            pass

    monkeypatch.setattr(pdfplumber, "open", lambda stream: _Pdf())
    with pytest.raises(ParserError, match="PDF_OCR_REQUIRED"):
        parse_document("application/pdf", b"mixed")


def test_docx_paragraphs_and_tables_are_extracted() -> None:
    document = docx.Document()
    document.add_paragraph("first paragraph iade")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "topic"
    table.rows[0].cells[1].text = "refund"
    buffer = io.BytesIO()
    document.save(buffer)

    out = parse_document(_DOCX_MIME, buffer.getvalue())
    assert "first paragraph iade" in out.text
    assert "topic | refund" in out.text
    assert out.parser == "docx"


def test_docx_malformed_fails_closed() -> None:
    with pytest.raises(ParserError) as exc:
        parse_document(_DOCX_MIME, b"not a docx zip")
    assert exc.value.code == "DOCX_PARSE_FAILED"


def test_xlsx_cells_become_pipe_joined_rows() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["topic", "detail"])
    worksheet.append(["refund", "iade window"])
    worksheet.append([None, None])  # blank row is skipped
    buffer = io.BytesIO()
    workbook.save(buffer)

    out = parse_document(_XLSX_MIME, buffer.getvalue())
    assert out.text == "topic | detail\nrefund | iade window"
    assert out.parser == "xlsx"
    assert out.element_count == 2


def test_xlsx_malformed_fails_closed() -> None:
    with pytest.raises(ParserError) as exc:
        parse_document(_XLSX_MIME, b"not an xlsx zip")
    assert exc.value.code == "XLSX_PARSE_FAILED"
